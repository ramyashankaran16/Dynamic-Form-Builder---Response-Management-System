"""Excel and PDF export of form responses."""
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.constants import ExportFormat, FieldType
from app.models import Form, FormResponse
from app.schemas.response import ResponseFilters
from app.services.form_engine import display_value
from app.services.form_service import get_form_or_404
from app.services.response_service import build_response_query, response_answers
from app.utils.http import slugify
from app.utils.time import utcnow

MEDIA_TYPES = {
    ExportFormat.EXCEL: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ExportFormat.PDF: "application/pdf",
}
EXTENSIONS = {ExportFormat.EXCEL: "xlsx", ExportFormat.PDF: "pdf"}
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
BRAND = colors.HexColor("#1F4E79")


def _load(db: Session, form_id: int, filters: ResponseFilters) -> tuple[Form, list[FormResponse]]:
    form = get_form_or_404(db, form_id)
    responses = list(db.scalars(build_response_query(form_id, filters)).unique().all())
    return form, responses


def _cell_value(field, value: Any) -> Any:
    if value is None:
        return ""
    if field.field_type in (FieldType.NUMBER, FieldType.RATING) and isinstance(value, (int, float)):
        return value
    return display_value(field.field_type, value, {o.value: o.label for o in field.options})


def build_excel(form: Form, responses: list[FormResponse], path: Path, filters: ResponseFilters) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Responses"
    fields = list(form.fields)
    headers = ["Response ID", "Submitted At (UTC)", "Submitted By", "Submitter Email", "Version", "Last Updated (UTC)"]
    headers += [f.label for f in fields]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for r in responses:
        answers = response_answers(r)
        ws.append(
            [
                r.id,
                r.created_at,
                r.submitter.full_name if r.submitter else "Anonymous",
                r.submitter.email if r.submitter else "",
                r.version,
                r.updated_at,
            ]
            + [_cell_value(f, answers.get(f.field_key)) for f in fields]
        )

    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for c in row:
            c.number_format = "yyyy-mm-dd hh:mm"
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for c in row:
            c.number_format = "yyyy-mm-dd hh:mm"
    for idx, column in enumerate(ws.columns, start=1):
        width = max((len(str(c.value)) if c.value is not None else 0) for c in column)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 12), 50)
    ws.freeze_panes = "A2"
    if responses:
        ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Summary")
    applied = {k: v for k, v in filters.model_dump(mode="json").items() if v not in (None, "", "all") and k not in ("sort_by", "sort_order")}
    rows = [
        ("Form", form.title),
        ("Description", form.description or ""),
        ("Status", "Active" if form.is_active else "Disabled"),
        ("Visibility", "Public" if form.is_public else "Authenticated users"),
        ("Owner", form.owner.full_name if form.owner else ""),
        ("Total responses exported", len(responses)),
        ("Filters applied", ", ".join(f"{k}={v}" for k, v in applied.items()) or "None"),
        ("Exported at (UTC)", utcnow().strftime("%Y-%m-%d %H:%M")),
    ]
    for key, value in rows:
        summary.append([key, value])
        summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 60
    wb.save(path)


def build_pdf(form: Form, responses: list[FormResponse], path: Path) -> None:
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=9, leading=11)
    cell_bold = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold")
    heading = ParagraphStyle("h", parent=styles["Heading3"], textColor=BRAND, spaceBefore=10, spaceAfter=4)

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(15 * mm, 10 * mm, f"{form.title[:80]} - responses export")
        canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=18 * mm,
        title=f"{form.title} - Responses",
    )
    story: list[Any] = [Paragraph(escape(form.title), styles["Title"])]
    if form.description:
        story.append(Paragraph(escape(form.description), styles["BodyText"]))
    meta = Table(
        [
            ["Status", "Active" if form.is_active else "Disabled", "Total responses", str(len(responses))],
            ["Visibility", "Public" if form.is_public else "Authenticated", "Exported at (UTC)", utcnow().strftime("%Y-%m-%d %H:%M")],
        ],
        colWidths=[28 * mm, 52 * mm, 35 * mm, 65 * mm],
    )
    meta.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F2F6FA")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#C9D6E3")),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [Spacer(1, 6), meta, Spacer(1, 6)]

    if not responses:
        story.append(Paragraph("No responses found for the selected filters.", styles["Italic"]))

    for r in responses:
        answers = response_answers(r)
        who = r.submitter.full_name if r.submitter else "Anonymous"
        story.append(Paragraph(
            f"Response #{r.id} &nbsp;|&nbsp; {escape(who)} &nbsp;|&nbsp; {r.created_at:%Y-%m-%d %H:%M} UTC"
            + (f" &nbsp;|&nbsp; v{r.version}" if r.version > 1 else ""),
            heading,
        ))
        data = [[Paragraph("Field", cell_bold), Paragraph("Answer", cell_bold)]]
        for f in form.fields:
            value = display_value(f.field_type, answers.get(f.field_key), {o.value: o.label for o in f.options})
            data.append([Paragraph(escape(f.label), cell), Paragraph(escape(value) or "-", cell)])
        t = Table(data, colWidths=[60 * mm, 120 * mm], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F0")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D6E3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ]))
        story.append(t)

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


def generate_export(db: Session, form_id: int, fmt: str, filters: ResponseFilters | dict | None = None) -> Path:
    """Create the export file in EXPORT_DIR and return its path."""
    fmt = ExportFormat(fmt)
    if not isinstance(filters, ResponseFilters):
        filters = ResponseFilters(**(filters or {}))
    form, responses = _load(db, form_id, filters)
    settings.export_dir.mkdir(parents=True, exist_ok=True)
    name = f"{slugify(form.title, 40)}_responses_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}.{EXTENSIONS[fmt]}"
    path = settings.export_dir / name
    if fmt == ExportFormat.EXCEL:
        build_excel(form, responses, path, filters)
    else:
        build_pdf(form, responses, path)
    return path
